#!/usr/bin/env python3
"""Builds a sample Google-Slides-style .pptx and Google-Docs-style .docx, runs the extractor, checks
the outline. Also writes the sample files for trying the import skill:
    python3 scripts/import/test_extract.py --keep imports/samples
"""
import json
import subprocess
import sys
import tempfile
from pathlib import Path

import docx
from docx.shared import Pt as DPt
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

HERE = Path(__file__).parent


def _png(w=4, h=4, rgb=(29, 74, 60)):
    import struct, zlib
    raw = b"".join(b"\x00" + bytes(rgb) * w for _ in range(h))
    chunk = lambda t, d: struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xffffffff)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


PNG = _png()


def make_pptx(path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Pine+ referral pilot: referred members stay longer"
    s.placeholders[1].text = "8-week pilot readout · Customer & Growth · Aug 2026"

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "The pilot beat its sign-up target by 38%"
    for i, (num, lab) in enumerate([("4,140", "referred sign-ups"), ("+38%", "vs. target of 3,000"), ("$21", "cost per sign-up")]):
        tb = s.shapes.add_textbox(Inches(0.5 + 3.1 * i), Inches(2), Inches(2.9), Inches(1.6)).text_frame
        tb.text = num
        tb.paragraphs[0].runs[0].font.size = Pt(40)
        p = tb.add_paragraph(); p.text = lab; p.runs[0].font.size = Pt(14)
    s.notes_slide.notes_text_frame.text = "Target was set in the Q2 plan. Cost per sign-up includes the $15 credit and ops time."

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Referred members churn less in months 2-4"
    cd = CategoryChartData()
    cd.categories = ["Month 1", "Month 2", "Month 3", "Month 4"]
    cd.add_series("Referred", (0.041, 0.052, 0.071, 0.049))
    cd.add_series("Paid social", (0.058, 0.083, 0.162, 0.079))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.6), Inches(9), Inches(4.5), cd)
    s.notes_slide.notes_text_frame.text = "Monthly-plan members only, n = 3,380 referred vs 5,120 paid social. Month 3 gap is significant (p < 0.01)."

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Results by channel"
    rows, cols = 4, 4
    t = s.shapes.add_table(rows, cols, Inches(0.5), Inches(1.6), Inches(9), Inches(2)).table
    data = [["Channel", "Sign-ups", "Month-3 churn", "Cost / sign-up"], ["Referral", "4,140", "7.1%", "$21"],
            ["Paid social", "5,120", "16.2%", "$34"], ["Search", "6,300", "11.8%", "$27"]]
    for r in range(rows):
        for c in range(cols):
            t.cell(r, c).text = data[r][c]

    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Recommendation: roll out to all members in October"
    tf = s.placeholders[1].text_frame
    tf.text = "Extend the referral credit to all Pine+ members from Oct 6"
    for txt, lvl in [("Owner: Membership squad", 1), ("Keep a 10% holdout to measure incrementality", 0), ("Revisit the $15 credit after 12 weeks", 0)]:
        p = tf.add_paragraph(); p.text = txt; p.level = lvl
    img = HERE / "_tmp.png"; img.write_bytes(PNG)
    s.shapes.add_picture(str(img), Inches(8), Inches(5), Inches(1), Inches(1)); img.unlink()
    prs.save(path)


def make_docx(path):
    d = docx.Document()
    d.add_paragraph("Returns policy test: 60-day window", style="Title")
    d.add_paragraph("What happened when we extended returns from 30 to 60 days", style="Subtitle")
    d.add_heading("Summary", level=1)
    p = d.add_paragraph("Extending the window lifted conversion by ")
    r = p.add_run("2.1%"); r.bold = True
    p.add_run(" with no increase in return rate.")
    d.add_paragraph("Conversion +2.1% (95% CI 0.8 to 3.4%)", style="List Bullet")
    d.add_paragraph("Return rate flat at 11.0%", style="List Bullet")
    d.add_heading("Results", level=1)
    t = d.add_table(rows=3, cols=3)
    for i, row in enumerate([["Metric", "30-day", "60-day"], ["Conversion", "3.40%", "3.47%"], ["Return rate", "11.0%", "11.0%"]]):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    d.add_heading("Next steps", level=1)
    d.add_paragraph("Ship the 60-day window to all categories", style="List Number")
    d.add_paragraph("Monitor furniture returns monthly", style="List Number")
    img = HERE / "_tmp.png"; img.write_bytes(PNG)
    d.add_picture(str(img)); img.unlink()
    d.save(path)


def run(src, out):
    subprocess.run([sys.executable, str(HERE / "extract.py"), str(src), "--out", str(out)], check=True)
    return json.loads((out / "outline.json").read_text())


def main():
    keep = sys.argv[sys.argv.index("--keep") + 1] if "--keep" in sys.argv else None
    base = Path(keep) if keep else Path(tempfile.mkdtemp())
    base.mkdir(parents=True, exist_ok=True)
    pptx_path, docx_path = base / "referral-pilot-deck.pptx", base / "returns-policy-doc.docx"
    make_pptx(pptx_path); make_docx(docx_path)

    o = run(pptx_path, base / "out-pptx")
    assert o["meta"]["source_type"] == "google-slides"
    assert o["meta"]["title"].startswith("Pine+ referral pilot"), o["meta"]
    assert len(o["slides"]) == 5
    assert o["slides"][1]["kpi_candidates"] == ["4,140", "+38%", "$21"], o["slides"][1]["kpi_candidates"]
    assert "Cost per sign-up" in o["slides"][1]["notes"]
    chart = next(b for b in o["slides"][2]["blocks"] if b["type"] == "chart")
    assert chart["pp_chart"]["type"] == "column" and chart["pp_chart"]["categories"][2] == "Month 3"
    assert chart["pp_chart"]["series"][1]["values"][2] == 0.162 and "_note" in chart["pp_chart"]
    table = next(b for b in o["slides"][3]["blocks"] if b["type"] == "table")
    assert table["rows"][2][2] == "16.2%"
    bullets = next(b for b in o["slides"][4]["blocks"] if b["type"] == "text")
    assert bullets["items"][1] == {**bullets["items"][1], "text": "Owner: Membership squad", "level": 1}
    assert any(b["type"] == "image" for b in o["slides"][4]["blocks"])

    o = run(docx_path, base / "out-docx")
    assert o["meta"]["title"] == "Returns policy test: 60-day window"
    assert o["meta"]["subtitle"].startswith("What happened")
    kinds = [b["type"] for b in o["blocks"]]
    assert kinds[:3] == ["heading", "paragraph", "list"], kinds
    assert "**2.1%**" in o["blocks"][1]["text"]
    lists = [b for b in o["blocks"] if b["type"] == "list"]
    assert lists[0]["ordered"] is False and lists[1]["ordered"] is True, lists
    assert any(b["type"] == "table" and b["rows"][1][2] == "3.47%" for b in o["blocks"])
    assert any(b["type"] == "image" for b in o["blocks"])
    print("import extractor: all checks passed" + (f" (samples in {base})" if keep else ""))


if __name__ == "__main__":
    main()
