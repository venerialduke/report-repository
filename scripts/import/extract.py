#!/usr/bin/env python3
"""
Extract a Google Doc or Google Slides deck into a structured outline the report-import skill can
turn into an Insights Hub report.

Input (any of):
  path/to/file.docx   Google Docs  -> File > Download > Microsoft Word (.docx)
  path/to/file.pptx   Google Slides -> File > Download > Microsoft PowerPoint (.pptx)
  path/to/file.md     Google Docs  -> File > Download > Markdown (.md)   (text only)
  --url <google link> Downloads the export directly; only works for files shared "Anyone with the link"

Output (in --out, default imports/<name>/):
  outline.json   ordered blocks (headings, paragraphs, lists, tables, images, charts, notes)
  outline.md     the same, human-readable, for review
  assets/        images extracted from the file

Usage:
  pip install -r scripts/import/requirements.txt
  python3 scripts/import/extract.py deck.pptx --out imports/referral-pilot
  python3 scripts/import/extract.py --url https://docs.google.com/presentation/d/<id>/edit
"""
import argparse
import json
import os
import re
import sys
import urllib.request
from pathlib import Path

try:
    import docx  # python-docx
    from docx.oxml.ns import qn
    from pptx import Presentation
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.util import Pt
except ImportError:  # pragma: no cover
    sys.exit("Missing dependencies. Run: pip install -r scripts/import/requirements.txt")


# ---------------------------------------------------------------- helpers
def slugify(s, n=60):
    s = re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")
    return s[:n].rstrip("-") or "imported-report"


class Assets:
    def __init__(self, out):
        self.dir = Path(out) / "assets"
        self.n = 0
        self.seen = {}

    def save(self, blob, ext, hint="image"):
        key = hash(blob)
        if key in self.seen:
            return self.seen[key]
        self.dir.mkdir(parents=True, exist_ok=True)
        self.n += 1
        ext = (ext or "png").lower().replace("jpeg", "jpg")
        name = f"{slugify(hint, 30)}-{self.n}.{ext}"
        (self.dir / name).write_bytes(blob)
        rel = f"assets/{name}"
        self.seen[key] = rel
        return rel


NUMBER = re.compile(r"[-+−]?\$?\d[\d,.]*\s?(%|pp|pts?|x|×|k|m|bn|b)?", re.I)


# ---------------------------------------------------------------- DOCX (Google Docs)
def _num_formats(document):
    """numId -> {ilvl: numFmt} so we can tell bullets from numbered lists."""
    out = {}
    try:
        numbering = document.part.numbering_part.element
    except Exception:
        return out
    abstract = {}
    for an in numbering.findall(qn("w:abstractNum")):
        aid = an.get(qn("w:abstractNumId"))
        lv = {}
        for l in an.findall(qn("w:lvl")):
            fmt = l.find(qn("w:numFmt"))
            lv[int(l.get(qn("w:ilvl")))] = fmt.get(qn("w:val")) if fmt is not None else "bullet"
        abstract[aid] = lv
    for num in numbering.findall(qn("w:num")):
        a = num.find(qn("w:abstractNumId"))
        if a is not None:
            out[num.get(qn("w:numId"))] = abstract.get(a.get(qn("w:val")), {})
    return out


def _numpr(p_el, style):
    """List numbering: set on the paragraph (Google Docs exports) or inherited from its style (Word)."""
    ppr = p_el.find(qn("w:pPr"))
    if ppr is not None and ppr.find(qn("w:numPr")) is not None:
        return ppr.find(qn("w:numPr"))
    while style is not None:
        sppr = style.element.find(qn("w:pPr"))
        if sppr is not None and sppr.find(qn("w:numPr")) is not None:
            return sppr.find(qn("w:numPr"))
        style = style.base_style
    return None


def _runs_md(paragraph):
    """Paragraph -> inline markdown (**bold**, *italic*, [link](url))."""
    parts = []
    for item in paragraph.iter_inner_content():
        if hasattr(item, "address"):  # Hyperlink
            text = item.text
            if text.strip():
                parts.append(f"[{text}]({item.address})" if item.address else text)
            continue
        text = item.text
        if not text:
            continue
        if item.bold and text.strip():
            text = f"**{text.strip()}**" + (" " if text.endswith(" ") else "")
        elif item.italic and text.strip():
            text = f"*{text.strip()}*" + (" " if text.endswith(" ") else "")
        parts.append(text)
    return re.sub(r"\*\*\s*\*\*", "", "".join(parts)).strip()


def _paragraph_images(paragraph, document, assets):
    imgs = []
    for blip in paragraph._p.iter(qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if not rid or rid not in document.part.related_parts:
            continue
        part = document.part.related_parts[rid]
        ext = part.partname.ext if hasattr(part.partname, "ext") else "png"
        alt = ""
        for d in paragraph._p.iter(qn("wp:docPr")):
            alt = d.get("descr") or d.get("title") or ""
        imgs.append({"type": "image", "src": assets.save(part.blob, ext), "alt": alt})
    return imgs


def extract_docx(path, assets):
    document = docx.Document(path)
    fmts = _num_formats(document)
    blocks, meta = [], {"source_type": "google-docs"}
    body = document.element.body
    lst = None
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "tbl":
            lst = None
            table = docx.table.Table(child, document)
            rows = [[" ".join(p.text for p in cell.paragraphs).strip() for cell in row.cells] for row in table.rows]
            if any(any(c for c in r) for r in rows):
                blocks.append({"type": "table", "rows": rows})
            continue
        if tag != "p":
            continue
        p = docx.text.paragraph.Paragraph(child, document)
        blocks.extend(_paragraph_images(p, document, assets))
        text = _runs_md(p)
        style = (p.style.name if p.style is not None else "") or ""
        numpr = _numpr(child, p.style)
        if not text:
            continue
        if style == "Title":
            meta.setdefault("title", p.text.strip()); lst = None; continue
        if style == "Subtitle":
            meta.setdefault("subtitle", p.text.strip()); lst = None; continue
        m = re.match(r"Heading (\d)", style)
        if m:
            lst = None
            blocks.append({"type": "heading", "level": int(m.group(1)), "text": p.text.strip()})
            continue
        if numpr is not None:
            ilvl_el, numid_el = numpr.find(qn("w:ilvl")), numpr.find(qn("w:numId"))
            ilvl = int(ilvl_el.get(qn("w:val"))) if ilvl_el is not None else 0
            numid = numid_el.get(qn("w:val")) if numid_el is not None else None
            fmt = fmts.get(numid, {}).get(ilvl)
            ordered = (fmt not in ("bullet", "none")) if fmt else ("Number" in style)
            if lst is None or lst["ordered"] != ordered:
                lst = {"type": "list", "ordered": ordered, "items": []}
                blocks.append(lst)
            lst["items"].append({"text": text, "level": ilvl})
            continue
        lst = None
        blocks.append({"type": "paragraph", "text": text})
    if "title" not in meta:
        h = next((b for b in blocks if b["type"] == "heading"), None)
        if h:
            meta["title"] = h["text"]
    meta["core_author"] = document.core_properties.author or ""
    return {"meta": meta, "blocks": blocks}


# ---------------------------------------------------------------- PPTX (Google Slides)
CHART_MAP = {
    "COLUMN_CLUSTERED": "column", "COLUMN_STACKED": "stacked", "COLUMN_STACKED_100": "stacked",
    "BAR_CLUSTERED": "bar", "BAR_STACKED": "stacked-bar", "BAR_STACKED_100": "stacked-bar",
    "LINE": "line", "LINE_MARKERS": "line", "LINE_STACKED": "line", "LINE_MARKERS_STACKED": "line",
    "PIE": "bar", "DOUGHNUT": "bar", "AREA": "line", "AREA_STACKED": "stacked",
}


def _chart(shape):
    ch = shape.chart
    try:
        ctype = XL_CHART_TYPE(ch.chart_type).name
    except Exception:
        ctype = str(ch.chart_type)
    plot = ch.plots[0]
    cats = [str(c) for c in plot.categories]
    series = [{"name": s.name or f"Series {i + 1}", "values": [None if v is None else round(float(v), 6) for v in s.values]}
              for i, s in enumerate(plot.series)]
    fmt = None
    try:
        if plot.has_data_labels and plot.data_labels.number_format not in ("General", ""):
            fmt = plot.data_labels.number_format
    except Exception:
        pass
    title = ""
    try:
        if ch.has_title and ch.chart_title.has_text_frame:
            title = ch.chart_title.text_frame.text.strip()
    except Exception:
        pass
    suggestion = {"type": CHART_MAP.get(ctype, "column"), "categories": cats, "series": series}
    if all(v is None or 0 <= v <= 1 for s in series for v in s["values"]) and series:
        suggestion["_note"] = "values look like fractions (0-1); multiply by 100 and use format pct/pct1"
    return {"type": "chart", "source_chart_type": ctype, "title": title, "number_format": fmt,
            "pp_chart": suggestion}


def _shape_blocks(shape, assets, slide_no):
    out = []
    st = shape.shape_type
    if st == MSO_SHAPE_TYPE.GROUP:
        for s in sorted(shape.shapes, key=lambda s: (s.top or 0, s.left or 0)):
            out.extend(_shape_blocks(s, assets, slide_no))
        return out
    if getattr(shape, "has_chart", False) and shape.has_chart:
        out.append(_chart(shape))
        return out
    if getattr(shape, "has_table", False) and shape.has_table:
        rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
        out.append({"type": "table", "rows": rows})
        return out
    if st == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "image"):
        try:
            img = shape.image
            out.append({"type": "image", "src": assets.save(img.blob, img.ext, f"slide-{slide_no}"),
                        "alt": (shape._element.xpath("./p:nvPicPr/p:cNvPr/@descr") or [""])[0]
                        , "note": "If this is a chart or table screenshot, ask the owner for the underlying data."})
        except Exception:
            pass
        return out
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        items, big = [], []
        for para in shape.text_frame.paragraphs:
            text = "".join(r.text for r in para.runs).strip()
            if not text:
                continue
            size = max([r.font.size.pt for r in para.runs if r.font.size is not None] or [0])
            bold = any(r.font.bold for r in para.runs)
            items.append({"text": text, "level": para.level, "size": size, "bold": bold})
            if size >= 28 and NUMBER.search(text) and len(text) <= 24:
                big.append(text)
        if items:
            out.append({"type": "text", "items": items, "big_numbers": big})
    return out


def extract_pptx(path, assets):
    prs = Presentation(path)
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        blocks = []
        for shape in sorted(slide.shapes, key=lambda s: (s.top or 0, s.left or 0)):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            blocks.extend(_shape_blocks(shape, assets, i))
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        kpis = [n for b in blocks if b["type"] == "text" for n in b.get("big_numbers", [])]
        slides.append({"n": i, "title": title, "blocks": blocks, "notes": notes, "kpi_candidates": kpis})
    meta = {"source_type": "google-slides"}
    if slides:
        meta["title"] = slides[0]["title"]
        first_text = next((b for b in slides[0]["blocks"] if b["type"] == "text"), None)
        if first_text:
            meta["subtitle"] = first_text["items"][0]["text"]
    meta["core_author"] = prs.core_properties.author or ""
    return {"meta": meta, "slides": slides}


# ---------------------------------------------------------------- Markdown
def extract_md(path, _assets):
    text = Path(path).read_text(encoding="utf-8")
    title = next((l[2:].strip() for l in text.splitlines() if l.startswith("# ")), "")
    return {"meta": {"source_type": "google-docs", "title": title}, "markdown": text}


# ---------------------------------------------------------------- download by link
def download(url, out_dir):
    m = re.search(r"docs\.google\.com/(document|presentation)/d/([A-Za-z0-9_-]+)", url)
    if not m:
        sys.exit("Not a Google Docs/Slides link. Expected docs.google.com/document/d/<id> or /presentation/d/<id>.")
    kind, fid = m.groups()
    export = (f"https://docs.google.com/document/d/{fid}/export?format=docx" if kind == "document"
              else f"https://docs.google.com/presentation/d/{fid}/export/pptx")
    ext = "docx" if kind == "document" else "pptx"
    req = urllib.request.Request(export, headers={"User-Agent": "insights-hub-import"})
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            ctype = r.headers.get("Content-Type", "")
            data = r.read()
    except Exception as e:
        sys.exit(f"Download failed ({e}). Export the file manually: File > Download > .{ext}")
    if "text/html" in ctype or data[:2] != b"PK":
        sys.exit("Google returned a sign-in page, so the file isn't shared 'Anyone with the link'. "
                 f"Export it manually (File > Download > .{ext}) or use a Google Drive connector.")
    out_dir.mkdir(parents=True, exist_ok=True)
    dest = out_dir / f"source.{ext}"
    dest.write_bytes(data)
    return dest, url


# ---------------------------------------------------------------- outline.md
def to_markdown(outline):
    lines = [f"# {outline['meta'].get('title', '(untitled)')}", ""]
    if outline["meta"].get("subtitle"):
        lines += [f"_{outline['meta']['subtitle']}_", ""]

    def block_md(b):
        t = b["type"]
        if t == "heading":
            return ["#" * min(6, b["level"] + 1) + " " + b["text"], ""]
        if t == "paragraph":
            return [b["text"], ""]
        if t == "list":
            return [("  " * it["level"]) + ("1. " if b["ordered"] else "- ") + it["text"] for it in b["items"]] + [""]
        if t == "text":
            return [("  " * it["level"]) + "- " + it["text"] + (f"  _(size {it['size']:.0f})_" if it["size"] >= 28 else "") for it in b["items"]] + [""]
        if t == "table":
            rows = b["rows"]
            if not rows:
                return []
            out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])]
            out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
            return out + [""]
        if t == "image":
            return [f"![{b.get('alt', '')}]({b['src']})", ""]
        if t == "chart":
            return [f"**[Native chart: {b['source_chart_type']}] {b['title']}**", "```json", json.dumps(b["pp_chart"], ensure_ascii=False), "```", ""]
        return []

    if "markdown" in outline:
        return outline["markdown"]
    if "blocks" in outline:
        for b in outline["blocks"]:
            lines += block_md(b)
    for s in outline.get("slides", []):
        lines += [f"## Slide {s['n']}: {s['title']}", ""]
        for b in s["blocks"]:
            lines += block_md(b)
        if s["kpi_candidates"]:
            lines += ["_KPI candidates: " + ", ".join(s["kpi_candidates"]) + "_", ""]
        if s["notes"]:
            lines += ["> Speaker notes: " + s["notes"].replace("\n", " "), ""]
    return "\n".join(lines)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file", nargs="?", help=".docx, .pptx or .md exported from Google Docs/Slides")
    ap.add_argument("--url", help="Google Docs/Slides link (must be shared 'Anyone with the link')")
    ap.add_argument("--out", help="output folder (default imports/<name>)")
    a = ap.parse_args()
    if not a.file and not a.url:
        ap.error("give a file or --url")

    source_url = None
    if a.url:
        out = Path(a.out or "imports/download")
        path, source_url = download(a.url, out)
    else:
        path = Path(a.file)
        if not path.exists():
            sys.exit(f"File not found: {path}")
        out = Path(a.out or f"imports/{slugify(path.stem)}")
    out.mkdir(parents=True, exist_ok=True)
    assets = Assets(out)

    ext = path.suffix.lower()
    if ext == ".docx":
        outline = extract_docx(path, assets)
    elif ext == ".pptx":
        outline = extract_pptx(path, assets)
    elif ext in (".md", ".markdown"):
        outline = extract_md(path, assets)
    else:
        sys.exit("Unsupported file type. Export Google Docs as .docx (or .md) and Google Slides as .pptx.")

    outline["meta"].update({"source_file": path.name, "source_url": source_url,
                            "suggested_id": slugify(outline["meta"].get("title") or path.stem)})
    (out / "outline.json").write_text(json.dumps(outline, indent=2, ensure_ascii=False), encoding="utf-8")
    (out / "outline.md").write_text(to_markdown(outline), encoding="utf-8")

    n_items = len(outline.get("slides", outline.get("blocks", [])))
    charts = sum(1 for s in outline.get("slides", [{"blocks": outline.get("blocks", [])}]) for b in s["blocks"] if b["type"] == "chart")
    imgs = assets.n
    print(f"Extracted {outline['meta']['source_type']} '{outline['meta'].get('title', '')}': "
          f"{n_items} {'slides' if 'slides' in outline else 'blocks'}, {charts} native charts, {imgs} images")
    print(f"  {out}/outline.json\n  {out}/outline.md" + (f"\n  {out}/assets/" if imgs else ""))


if __name__ == "__main__":
    main()
